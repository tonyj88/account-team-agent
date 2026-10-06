"""Render plans into the confidential Account Plan template."""

from __future__ import annotations

import re
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.table import Table, _Cell

from atb.catalog import Catalog, FieldSpec, RowSectionSpec, load_catalog
from atb.plan import Evidence, Plan, PlanRow, PlanValue


@dataclass
class RenderReport:
    """Summary of values placed in a rendered document."""

    filled: list[str]
    missing_labels: list[str]
    rows_written: dict[str, int]


def _normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().removesuffix(":").strip().casefold()


def _cell_text(cell: _Cell) -> str:
    return " ".join(p.text for p in cell.paragraphs).strip()


def _cells(table: Table):
    for row in table.rows:
        for cell in row.cells:
            yield cell
            for nested in cell.tables:
                yield from _cells(nested)


def _tables(document) -> list[Table]:
    found: list[Table] = []

    def visit(table: Table) -> None:
        found.append(table)
        for cell in _cells_shallow(table):
            for nested in cell.tables:
                visit(nested)

    for table in document.tables:
        visit(table)
    return found


def _cells_shallow(table: Table):
    for row in table.rows:
        yield from row.cells


def _distinct(cells) -> list[_Cell]:
    """Row cells with horizontally merged duplicates removed (python-docx repeats them)."""
    seen: set[int] = set()
    result: list[_Cell] = []
    for cell in cells:
        if id(cell._tc) not in seen:
            seen.add(id(cell._tc))
            result.append(cell)
    return result


def _set_cell(cell: _Cell, value: str) -> None:
    paragraphs = cell.paragraphs
    first = paragraphs[0]
    runs = first.runs
    if runs:
        runs[0].text = value
        for run in runs[1:]:
            run.text = ""
    else:
        first.add_run(value)
    for paragraph in paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)


def _marker(value: str, item: PlanValue | PlanRow) -> str:
    if isinstance(item, PlanValue) and item.flag:
        return f"{value} ⚠ ({item.flag})"
    if item.status == "needs_approval":
        return f"{value} ⚠"
    return value


def _ordered_fields(catalog: Catalog) -> list[FieldSpec]:
    return list(catalog.fields.values())


def _row_table(tables: list[Table], section: RowSectionSpec) -> Table | None:
    labels = [_normalize(column.label) for column in section.columns]
    for table in tables:
        if not table.rows:
            continue
        header = [_normalize(_cell_text(cell)) for cell in table.rows[0].cells]
        cursor = 0
        for label in labels:
            try:
                cursor = header.index(label, cursor) + 1
            except ValueError:
                break
        else:
            return table
    return None


def _clear_row(row) -> None:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.text = ""


def _write_rows(table: Table, section: RowSectionSpec, rows: list[PlanRow]) -> int:
    positions: list[int] = []
    header = [_normalize(_cell_text(cell)) for cell in table.rows[0].cells]
    cursor = 0
    for column in section.columns:
        cursor = header.index(_normalize(column.label), cursor)
        positions.append(cursor)
        cursor += 1
    count = 0
    for row_data in rows:
        target = None
        for candidate in table.rows[1:]:
            if all(not _cell_text(cell) for cell in candidate.cells):
                target = candidate
                break
        if target is None:
            copied = deepcopy(table.rows[-1]._tr)
            table._tbl.append(copied)
            target = table.rows[-1]
            _clear_row(target)
        for index, column in zip(positions, section.columns, strict=True):
            value = row_data.cells.get(column.key, "")
            if column == section.columns[0] and row_data.status == "needs_approval":
                value += " ⚠"
            _set_cell(target.cells[index], value)
        count += 1
    return count


def render_docx(
    plan: Plan, template: Path, out: Path, catalog: Catalog | None = None
) -> RenderReport:
    """Fill catalog-labeled cells and row tables in a Word document."""
    catalog = catalog or load_catalog()
    document = Document(template)
    tables = _tables(document)
    all_cells = [
        (table, cells, index, cell)
        for table in tables
        for row in table.rows
        for cells in [_distinct(row.cells)]
        for index, cell in enumerate(cells)
    ]
    filled: list[str] = []
    missing: list[str] = []
    duplicate_labels: dict[str, list[_Cell]] = {}
    for spec in _ordered_fields(catalog):
        item = plan.fields.get(spec.key)
        if item is None or item.value is None:
            continue
        matches = [
            entry
            for entry in all_cells
            if _normalize(_cell_text(entry[3])) == _normalize(spec.label)
        ]
        if spec.section == "meddpicc":
            duplicate_labels.setdefault(spec.label, matches)
            matches = duplicate_labels[spec.label]
            index = 1 if spec.key.endswith(".status_rag") else 0
        else:
            index = 0
        if not matches:
            missing.append(spec.label)
            continue
        _, row_cells, cell_index, _ = matches[0]
        target_index = cell_index + index + 1
        if target_index >= len(row_cells):
            missing.append(spec.label)
            continue
        _set_cell(row_cells[target_index], _marker(item.value, item))
        filled.append(spec.key)

    rows_written: dict[str, int] = {}
    for section in catalog.row_sections.values():
        rows = plan.rows.get(section.key, [])
        if not rows:
            continue
        table = _row_table(tables, section)
        if table is None:
            missing.extend(column.label for column in section.columns)
            rows_written[section.key] = 0
            continue
        rows_written[section.key] = _write_rows(table, section, rows)
    document.save(out)
    return RenderReport(
        filled=filled, missing_labels=list(dict.fromkeys(missing)), rows_written=rows_written
    )


def _evidence_lines(sources: list[Evidence]) -> list[str]:
    result: list[str] = []
    for source in sources:
        result.append(f"- {source.kind}, {source.as_of}, {source.ref}")
        if source.quote:
            attribution = " — ".join(part for part in (source.speaker, source.timestamp) if part)
            result.append(f"  > {source.quote}" + (f" — {attribution}" if attribution else ""))
    return result


def evidence_markdown(plan: Plan, catalog: Catalog | None = None) -> str:
    """Build a Markdown evidence record in catalog order."""
    catalog = catalog or load_catalog()
    lines = [f"# {plan.account} — evidence ({plan.plan_date})", ""]
    for spec in _ordered_fields(catalog):
        value = plan.fields.get(spec.key)
        if value is None or value.value is None:
            continue
        lines.extend([f"## {spec.label}", "", value.value, "", f"Status: {value.status}"])
        if value.reasons:
            lines.append("Reasons: " + "; ".join(value.reasons))
        lines.extend(_evidence_lines(value.sources))
        if value.history:
            lines.extend(["", "History"])
            for old in value.history:
                lines.append(f"- {old.value or ''} ({old.evidence.as_of}, {old.evidence.kind})")
        lines.append("")
    for section in catalog.row_sections.values():
        rows = plan.rows.get(section.key, [])
        if not rows:
            continue
        lines.extend([f"## {section.title}", ""])
        for row in rows:
            lines.append(f"### {row.row_key}")
            lines.append(f"Status: {row.status}")
            if row.reasons:
                lines.append("Reasons: " + "; ".join(row.reasons))
            for column in section.columns:
                if column.key in row.cells:
                    lines.append(f"- **{column.label}:** {row.cells[column.key]}")
            lines.extend(_evidence_lines(row.sources))
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"
