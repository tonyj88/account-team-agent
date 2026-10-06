from datetime import date

from docx import Document

from atb.catalog import Catalog, ColumnSpec, FieldSpec, RowSectionSpec
from atb.plan import Candidate, Evidence, Plan, PlanRow, PlanValue
from atb.render import evidence_markdown, render_docx


def _catalog() -> Catalog:
    fields = {
        "snapshot.account": FieldSpec(
            "snapshot.account", "Account", "snapshot", "system", "auto", True
        ),
        "snapshot.account_owner": FieldSpec(
            "snapshot.account_owner", "Account Owner", "snapshot", "crm", "auto", True
        ),
        "snapshot.current_arr": FieldSpec(
            "snapshot.current_arr", "Current ARR", "snapshot", "crm", "auto", True
        ),
        "snapshot.renewal_dates": FieldSpec(
            "snapshot.renewal_dates", "Renewal Date(s)", "snapshot", "crm", "auto", True
        ),
        "meddpicc.metrics.evidence": FieldSpec(
            "meddpicc.metrics.evidence", "Metrics", "meddpicc", "note", "auto", True
        ),
        "meddpicc.metrics.status_rag": FieldSpec(
            "meddpicc.metrics.status_rag",
            "Metrics",
            "meddpicc",
            "llm_draft",
            "needs_approval",
            True,
        ),
        "absent.field": FieldSpec("absent.field", "Missing Field", "snapshot", "crm", "auto", True),
    }
    columns = tuple(
        ColumnSpec(key, label, "note", "auto")
        for key, label in [
            ("name", "Name"),
            ("role", "Role / Title"),
            ("meddpicc_role", "MEDDPICC role"),
            ("position", "Position"),
            ("coverage", "Coverage"),
        ]
    )
    return Catalog(
        fields, {"stakeholders": RowSectionSpec("stakeholders", "Stakeholders", True, columns)}
    )


def test_render_docx_labels_rows_and_markers(tmp_path):
    template = tmp_path / "template.docx"
    output = tmp_path / "output.docx"
    doc = Document()
    snapshot = doc.add_table(rows=4, cols=2)
    for row, label in zip(
        snapshot.rows, ["Account", "Account Owner", "Current ARR", "Renewal Date(s)"], strict=True
    ):
        row.cells[0].text = label
    meddpicc = doc.add_table(rows=3, cols=3)
    meddpicc.rows[0].cells[0].text = "label"
    meddpicc.rows[0].cells[1].text = "evidence"
    meddpicc.rows[0].cells[2].text = "R/A/G"
    for row, label in zip(meddpicc.rows[1:], ["Metrics", "Champion"], strict=True):
        row.cells[0].text = label
    stakeholders = doc.add_table(rows=2, cols=5)
    for cell, label in zip(
        stakeholders.rows[0].cells,
        ["Name", "Role / Title", "MEDDPICC role", "Position", "Coverage"],
        strict=True,
    ):
        cell.text = label
    doc.save(template)

    source = Evidence(
        kind="note",
        as_of=date(2026, 10, 1),
        ref="notes/call.md",
        quote="We will expand by 20%.",
        speaker="Jane Doe",
        timestamp="00:01:02",
    )
    plan = Plan(
        account="Acme Corp",
        plan_date=date(2026, 10, 6),
        fields={
            "snapshot.account": PlanValue(value="Acme Corp", status="auto"),
            "snapshot.account_owner": PlanValue(value="Jane Doe", status="needs_approval"),
            "snapshot.current_arr": PlanValue(
                value="$100K", status="auto", flag="customer says $120K"
            ),
            "snapshot.renewal_dates": PlanValue(value="2027-01-01", status="auto"),
            "meddpicc.metrics.evidence": PlanValue(
                value="20% expansion",
                status="auto",
                sources=[source],
                history=[
                    Candidate(
                        field="meddpicc.metrics.evidence",
                        value="10% growth",
                        confidence="medium",
                        evidence=Evidence(
                            kind="baseline", as_of=date(2025, 10, 1), ref="baseline.docx"
                        ),
                    )
                ],
            ),
            "meddpicc.metrics.status_rag": PlanValue(value="Amber", status="needs_approval"),
            "absent.field": PlanValue(value="not placed", status="auto"),
        },
        rows={
            "stakeholders": [
                PlanRow(
                    row_key=f"person-{i}",
                    cells={
                        "name": f"Person {i}",
                        "role": "Buyer",
                        "meddpicc_role": "Champion",
                        "position": "Supporter",
                        "coverage": "Monthly",
                    },
                    status="needs_approval" if i == 1 else "auto",
                )
                for i in range(1, 4)
            ]
        },
    )
    report = render_docx(plan, template, output, _catalog())
    rendered = Document(output)
    assert rendered.tables[0].rows[0].cells[1].text == "Acme Corp"
    assert rendered.tables[0].rows[1].cells[1].text == "Jane Doe ⚠"
    assert rendered.tables[0].rows[2].cells[1].text == "$100K ⚠ (customer says $120K)"
    assert rendered.tables[1].rows[1].cells[1].text == "20% expansion"
    assert rendered.tables[1].rows[1].cells[2].text == "Amber ⚠"
    assert len(rendered.tables[2].rows) == 4
    assert rendered.tables[2].rows[1].cells[0].text == "Person 1 ⚠"
    assert "Missing Field" in report.missing_labels
    assert report.rows_written["stakeholders"] == 3

    evidence = evidence_markdown(plan, _catalog())
    assert "> We will expand by 20%. — Jane Doe — 00:01:02" in evidence
    assert "10% growth (2025-10-01, baseline)" in evidence


def test_render_skips_merged_label_cells(tmp_path):
    template = tmp_path / "merged.docx"
    document = Document()
    table = document.add_table(rows=1, cols=3)
    label = table.cell(0, 0).merge(table.cell(0, 1))
    label.text = "Account Owner"
    document.save(template)
    plan = Plan(
        account="Acme Corp",
        plan_date=date(2026, 10, 1),
        fields={
            "snapshot.account_owner": PlanValue(
                value="Jane Doe",
                status="auto",
                sources=[
                    Evidence(
                        kind="salesforce",
                        as_of=date(2026, 9, 1),
                        ref="Account/001000000000001AAA/OwnerId",
                    )
                ],
            )
        },
    )
    out = tmp_path / "out.docx"
    report = render_docx(plan, template, out, _catalog())
    row = Document(out).tables[0].rows[0]
    assert report.filled == ["snapshot.account_owner"]
    assert row.cells[0].text == "Account Owner"
    assert row.cells[2].text == "Jane Doe"
